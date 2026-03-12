"""Run a full legacy CLI parity matrix against oracle and Python CLI."""

from __future__ import annotations

import argparse
import json
import os
import shutil
import stat
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any

from parity_common import (
    CODEX_89DB_DIR,
    LEGACY_ORACLE_EXE,
    ORIGINAL_DIR,
    SNAPSHOT_DIR,
    list_mp3_files,
    parse_mp3gain_output,
    sha256_file,
)

if TYPE_CHECKING:
    from argparse import Namespace


DEFAULT_JOBS = 8
PYTHON_CLI_MODULE = "mp3gain_gui_py.legacy_cli"
MUTATING_SETUP_ARGS = ("/q", "/r", "/c")


@dataclass(frozen=True)
class _Case:
    name: str
    args: tuple[str, ...]
    mutates: bool
    gating: bool
    table_case: bool = False
    timestamp_case: bool = False
    uses_files: bool = True
    setup_args: tuple[str, ...] = ()


CASES: tuple[_Case, ...] = (
    _Case(
        name="info_version", args=("/v",), mutates=False, gating=False, uses_files=False
    ),
    _Case(
        name="info_help_h", args=("/h",), mutates=False, gating=False, uses_files=False
    ),
    _Case(
        name="info_help_qmark",
        args=("/?",),
        mutates=False,
        gating=False,
        uses_files=False,
    ),
    _Case(
        name="info_help_wrap",
        args=("/?", "wrap"),
        mutates=False,
        gating=False,
        uses_files=False,
    ),
    _Case(
        name="scan_qo", args=("/q", "/o"), mutates=False, gating=True, table_case=True
    ),
    _Case(
        name="scan_qo_sc",
        args=("/q", "/o", "/s", "c"),
        mutates=False,
        gating=True,
        table_case=True,
    ),
    _Case(
        name="scan_qo_x",
        args=("/q", "/o", "/x"),
        mutates=False,
        gating=True,
        table_case=True,
    ),
    _Case(
        name="scan_qo_sr",
        args=("/q", "/o", "/s", "r"),
        mutates=False,
        gating=True,
        table_case=True,
    ),
    _Case(
        name="scan_qo_ss",
        args=("/q", "/o", "/s", "s"),
        mutates=False,
        gating=True,
        table_case=True,
    ),
    _Case(
        name="scan_qo_e",
        args=("/q", "/o", "/e"),
        mutates=False,
        gating=True,
        table_case=True,
    ),
    _Case(name="track_rc", args=("/q", "/r", "/c"), mutates=True, gating=True),
    _Case(
        name="track_rc_ss",
        args=("/q", "/r", "/c", "/s", "s"),
        mutates=True,
        gating=True,
    ),
    _Case(
        name="track_rc_sr",
        args=("/q", "/r", "/c", "/s", "r"),
        mutates=True,
        gating=True,
    ),
    _Case(
        name="track_rc_d15",
        args=("/q", "/r", "/c", "/d", "1.5"),
        mutates=True,
        gating=True,
    ),
    _Case(
        name="track_rc_m1",
        args=("/q", "/r", "/c", "/m", "1"),
        mutates=True,
        gating=True,
    ),
    _Case(name="track_rc_w", args=("/q", "/r", "/c", "/w"), mutates=True, gating=True),
    _Case(name="track_rk", args=("/q", "/r", "/k"), mutates=True, gating=True),
    _Case(name="track_rc_t", args=("/q", "/r", "/c", "/t"), mutates=True, gating=True),
    _Case(
        name="track_rc_p",
        args=("/q", "/r", "/c", "/p"),
        mutates=True,
        gating=True,
        timestamp_case=True,
    ),
    _Case(name="track_rc_f", args=("/q", "/r", "/c", "/f"), mutates=True, gating=True),
    _Case(name="album_ac", args=("/q", "/a", "/c"), mutates=True, gating=True),
    _Case(name="album_ack", args=("/q", "/a", "/c", "/k"), mutates=True, gating=True),
    _Case(
        name="album_ac_mneg1_dneg15",
        args=("/q", "/a", "/c", "/m", "-1", "/d", "-1.5"),
        mutates=True,
        gating=True,
    ),
    _Case(name="direct_g4", args=("/q", "/g", "4"), mutates=True, gating=True),
    _Case(name="direct_l0_3", args=("/q", "/l", "0", "3"), mutates=True, gating=True),
    _Case(
        name="direct_l1_neg3", args=("/q", "/l", "1", "-3"), mutates=True, gating=True
    ),
    _Case(
        name="undo_after_apply",
        args=("/q", "/u"),
        mutates=True,
        gating=True,
        setup_args=MUTATING_SETUP_ARGS,
    ),
    _Case(
        name="tags_sa_sd", args=("/q", "/s", "a", "/s", "d"), mutates=True, gating=True
    ),
    _Case(
        name="tags_si_sd", args=("/q", "/s", "i", "/s", "d"), mutates=True, gating=True
    ),
)


def _parse_args() -> Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=ORIGINAL_DIR,
        help="Source directory with original MP3 files.",
    )
    parser.add_argument(
        "--work-root",
        type=Path,
        default=CODEX_89DB_DIR / "cli_option_parity",
        help="Root workspace for per-run isolated command copies.",
    )
    parser.add_argument(
        "--snapshot-dir",
        type=Path,
        default=SNAPSHOT_DIR,
        help="Directory for JSON/JSONL/XLSX artifacts.",
    )
    parser.add_argument(
        "--oracle-exe",
        type=Path,
        default=LEGACY_ORACLE_EXE,
        help="Path to legacy mp3gain.exe oracle binary.",
    )
    parser.add_argument(
        "--jobs",
        type=int,
        default=DEFAULT_JOBS,
        help="Parallel worker count (default 8, capped by workload).",
    )
    parser.add_argument(
        "--sample",
        type=int,
        default=0,
        help="If > 0, limit to first N source files.",
    )
    return parser.parse_args()


def _timestamp_slug() -> str:
    return datetime.now(tz=UTC).strftime("%Y%m%d_%H%M%S")


def _make_writable(path: Path) -> None:
    os.chmod(path, stat.S_IREAD | stat.S_IWRITE)


def _copy_writable(src: Path, dst: Path) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dst)
    _make_writable(dst)


def _run_command(command: list[str], *, cwd: Path) -> dict[str, Any]:
    started = time.perf_counter()
    completed = subprocess.run(
        command, cwd=cwd, check=False, capture_output=True, text=True
    )
    return {
        "command": command,
        "exit_code": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "duration_s": round(time.perf_counter() - started, 6),
    }


def _run_pair(
    *,
    oracle_command: list[str],
    python_command: list[str],
    cwd: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    with ThreadPoolExecutor(max_workers=2) as executor:
        oracle_future = executor.submit(_run_command, oracle_command, cwd=cwd)
        python_future = executor.submit(_run_command, python_command, cwd=cwd)
        return oracle_future.result(), python_future.result()


def _normalize_table(text: str) -> dict[str, dict[str, str]]:
    rows = parse_mp3gain_output(text)
    norm: dict[str, dict[str, str]] = {}
    for key, fields in rows.items():
        norm[key] = {field: value.strip() for field, value in fields.items()}
    return norm


def _compare_table(oracle_text: str, python_text: str) -> tuple[bool, dict[str, Any]]:
    oracle_rows = _normalize_table(oracle_text)
    python_rows = _normalize_table(python_text)
    missing_in_python = sorted(key for key in oracle_rows if key not in python_rows)
    missing_in_oracle = sorted(key for key in python_rows if key not in oracle_rows)
    value_diffs: dict[str, dict[str, dict[str, str]]] = {}
    for key in sorted(set(oracle_rows) & set(python_rows)):
        diffs: dict[str, dict[str, str]] = {}
        for field in sorted(set(oracle_rows[key]) | set(python_rows[key])):
            o_value = oracle_rows[key].get(field)
            p_value = python_rows[key].get(field)
            if o_value != p_value:
                diffs[field] = {"oracle": str(o_value), "python": str(p_value)}
        if diffs:
            value_diffs[key] = diffs
    ok = not missing_in_python and not missing_in_oracle and not value_diffs
    return ok, {
        "missing_in_python": missing_in_python,
        "missing_in_oracle": missing_in_oracle,
        "value_diffs": value_diffs,
    }


def _apply_c_for_mutation(case: _Case) -> tuple[str, ...]:
    if case.mutates and "/c" not in case.args:
        return (*case.args, "/c")
    return case.args


def _file_task_id(case: _Case, source_file: Path | None) -> str:
    if source_file is None:
        return case.name
    return f"{case.name}:{source_file.name}"


def _require_runtime_dependencies(repo_root: Path) -> None:
    hatch_check = subprocess.run(
        ["hatch", "--version"],
        cwd=repo_root,
        check=False,
        capture_output=True,
        text=True,
    )
    if hatch_check.returncode != 0:
        raise RuntimeError(
            "Missing 'hatch' executable in PATH; parity harness requires hatch run."
        )

    dep_check = subprocess.run(
        [
            "hatch",
            "run",
            "python",
            "-c",
            "import miniaudio, mutagen, openpyxl",
        ],
        cwd=repo_root,
        check=False,
        capture_output=True,
        text=True,
    )
    if dep_check.returncode != 0:
        stderr = (dep_check.stderr or dep_check.stdout).strip()
        raise RuntimeError(
            "Hatch runtime dependencies unavailable (expected miniaudio/mutagen/openpyxl). "
            f"Details: {stderr}"
        )


def _run_case_for_file(
    *,
    repo_root: Path,
    run_root: Path,
    oracle_exe: Path,
    case: _Case,
    source_file: Path | None,
) -> dict[str, Any]:
    case_args = _apply_c_for_mutation(case)
    task_id = _file_task_id(case, source_file)

    oracle_file: Path | None = None
    python_file: Path | None = None
    input_sha256: str | None = None
    pre_mtime_oracle: float | None = None
    pre_mtime_python: float | None = None

    if source_file is not None:
        task_root = run_root / case.name / source_file.name
        oracle_file = task_root / "oracle" / source_file.name
        python_file = task_root / "python" / source_file.name
        _copy_writable(source_file, oracle_file)
        _copy_writable(source_file, python_file)
        input_sha256 = sha256_file(source_file)
        pre_mtime_oracle = oracle_file.stat().st_mtime
        pre_mtime_python = python_file.stat().st_mtime

    oracle_target = str(oracle_file) if oracle_file is not None else None
    python_target = str(python_file) if python_file is not None else None

    oracle_setup: dict[str, Any] | None = None
    python_setup: dict[str, Any] | None = None
    if case.setup_args and oracle_target is not None and python_target is not None:
        oracle_setup_cmd = [str(oracle_exe), *case.setup_args, oracle_target]
        python_setup_cmd = [
            "hatch",
            "run",
            "python",
            "-m",
            PYTHON_CLI_MODULE,
            *case.setup_args,
            python_target,
        ]
        oracle_setup, python_setup = _run_pair(
            oracle_command=oracle_setup_cmd,
            python_command=python_setup_cmd,
            cwd=repo_root,
        )

    oracle_command = [str(oracle_exe), *case_args]
    python_command = ["hatch", "run", "python", "-m", PYTHON_CLI_MODULE, *case_args]
    if oracle_target is not None:
        oracle_command.append(oracle_target)
    if python_target is not None:
        python_command.append(python_target)

    oracle_run, python_run = _run_pair(
        oracle_command=oracle_command,
        python_command=python_command,
        cwd=repo_root,
    )

    oracle_sha = sha256_file(oracle_file) if oracle_file is not None else None
    python_sha = sha256_file(python_file) if python_file is not None else None
    post_mtime_oracle = oracle_file.stat().st_mtime if oracle_file is not None else None
    post_mtime_python = python_file.stat().st_mtime if python_file is not None else None

    exit_code_match = oracle_run["exit_code"] == python_run["exit_code"]
    byte_identity = oracle_sha == python_sha if oracle_sha is not None else True

    table_match = True
    table_diff: dict[str, Any] = {}
    if case.table_case:
        table_match, table_diff = _compare_table(
            oracle_run["stdout"], python_run["stdout"]
        )

    non_mutating_ok = True
    if source_file is not None and not case.mutates and input_sha256 is not None:
        non_mutating_ok = oracle_sha == input_sha256 and python_sha == input_sha256

    timestamp_ok = True
    timestamp_detail: dict[str, Any] = {}
    if case.timestamp_case and source_file is not None:
        oracle_preserved = (
            pre_mtime_oracle is not None
            and post_mtime_oracle is not None
            and abs(post_mtime_oracle - pre_mtime_oracle) <= 1.0
        )
        python_preserved = (
            pre_mtime_python is not None
            and post_mtime_python is not None
            and abs(post_mtime_python - pre_mtime_python) <= 1.0
        )
        timestamp_ok = oracle_preserved and python_preserved
        timestamp_detail = {
            "oracle_pre": pre_mtime_oracle,
            "oracle_post": post_mtime_oracle,
            "python_pre": pre_mtime_python,
            "python_post": post_mtime_python,
            "oracle_preserved": oracle_preserved,
            "python_preserved": python_preserved,
        }

    setup_ok = True
    if oracle_setup is not None and python_setup is not None:
        setup_ok = oracle_setup["exit_code"] == python_setup["exit_code"] == 0

    gate_checks: list[bool] = [setup_ok, exit_code_match]
    if case.mutates and source_file is not None:
        gate_checks.append(byte_identity)
    if not case.mutates and source_file is not None:
        gate_checks.append(non_mutating_ok)
    if case.table_case:
        gate_checks.append(table_match)
    if case.timestamp_case:
        gate_checks.append(timestamp_ok)

    case_pass = all(gate_checks) if case.gating else exit_code_match

    return {
        "task_id": task_id,
        "case": case.name,
        "file": source_file.name if source_file is not None else None,
        "gating": case.gating,
        "mutates": case.mutates,
        "setup_ok": setup_ok,
        "exit_code_match": exit_code_match,
        "byte_identity": byte_identity,
        "non_mutating_ok": non_mutating_ok,
        "table_match": table_match,
        "timestamp_ok": timestamp_ok,
        "case_pass": case_pass,
        "gate_checks": {
            "setup_ok": setup_ok,
            "exit_code_match": exit_code_match,
            "byte_identity": byte_identity,
            "non_mutating_ok": non_mutating_ok,
            "table_match": table_match,
            "timestamp_ok": timestamp_ok,
        },
        "table_diff": table_diff,
        "timestamp_detail": timestamp_detail,
        "input_sha256": input_sha256,
        "oracle_sha256": oracle_sha,
        "python_sha256": python_sha,
        "oracle": oracle_run,
        "python": python_run,
        "oracle_setup": oracle_setup,
        "python_setup": python_setup,
    }


def _write_progress_event(handle: Any, event: dict[str, Any]) -> None:
    handle.write(json.dumps(event, sort_keys=True) + "\n")
    handle.flush()


def _eta_text(*, completed: int, total: int, started: float) -> str:
    if completed <= 0:
        return "?"
    elapsed = max(0.001, time.perf_counter() - started)
    remaining = max(0, total - completed)
    eta = elapsed / float(completed) * float(remaining)
    return f"{eta:.1f}s"


def _write_xlsx(
    *,
    summary: dict[str, Any],
    case_rows: list[dict[str, Any]],
    per_file_rows: list[dict[str, Any]],
    failure_rows: list[dict[str, Any]],
    path: Path,
) -> None:
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Key", "Value"])
    for key in sorted(summary):
        ws_summary.append([key, str(summary[key])])

    ws_cases = wb.create_sheet("Cases")
    ws_cases.append(["Case", "Runs", "Pass", "Fail", "Gating"])
    for row in case_rows:
        ws_cases.append(
            [row["case"], row["runs"], row["pass"], row["fail"], row["gating"]]
        )

    ws_per_file = wb.create_sheet("PerFile")
    ws_per_file.append(
        [
            "Task",
            "Case",
            "File",
            "Gating",
            "Mutates",
            "Pass",
            "ExitMatch",
            "ByteIdentity",
            "NonMutatingOk",
            "TableMatch",
            "TimestampOk",
        ]
    )
    for row in per_file_rows:
        ws_per_file.append(
            [
                row["task_id"],
                row["case"],
                row["file"] or "",
                row["gating"],
                row["mutates"],
                row["case_pass"],
                row["exit_code_match"],
                row["byte_identity"],
                row["non_mutating_ok"],
                row["table_match"],
                row["timestamp_ok"],
            ]
        )

    ws_fail = wb.create_sheet("Failures")
    ws_fail.append(["Task", "Case", "File", "Details"])
    for row in failure_rows:
        ws_fail.append(
            [
                row["task_id"],
                row["case"],
                row["file"] or "",
                json.dumps(row["gate_checks"], sort_keys=True),
            ]
        )

    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    bold = Font(bold=True)
    for ws in (ws_summary, ws_cases, ws_per_file, ws_fail):
        for cell in ws[1]:
            cell.font = bold

    for ws in (ws_cases, ws_per_file):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            pass_cell = row[2] if ws is ws_cases else row[5]
            pass_value = str(pass_cell.value).strip().lower()
            if pass_value in {"true", "pass", "1"}:
                pass_cell.fill = green
            elif pass_value in {"false", "fail", "0"}:
                pass_cell.fill = red

    wb.save(path)


def main() -> int:
    args = _parse_args()
    repo_root = Path(__file__).resolve().parents[1]
    if not args.oracle_exe.exists():
        raise RuntimeError(f"Missing legacy oracle binary: {args.oracle_exe}")
    if not args.source_dir.exists():
        raise RuntimeError(f"Missing source directory: {args.source_dir}")

    _require_runtime_dependencies(repo_root)

    files = list_mp3_files(args.source_dir)
    if args.sample > 0:
        files = files[: args.sample]
    if not files:
        raise RuntimeError(f"No MP3 files found under {args.source_dir}")

    run_slug = _timestamp_slug()
    run_root = args.work_root / f"run_{run_slug}"
    run_root.mkdir(parents=True, exist_ok=True)
    args.snapshot_dir.mkdir(parents=True, exist_ok=True)

    source_hash_before = {path.name: sha256_file(path) for path in files}
    total_tasks = sum(len(files) for case in CASES if case.uses_files) + sum(
        1 for case in CASES if not case.uses_files
    )
    jobs = max(1, min(args.jobs, total_tasks))

    progress_path = args.snapshot_dir / f"parity_cli_full_progress_{run_slug}.jsonl"
    results_path = args.snapshot_dir / f"parity_cli_full_results_{run_slug}.json"
    run_path = args.snapshot_dir / f"parity_cli_full_run_{run_slug}.json"
    report_xlsx = args.snapshot_dir / f"parity_cli_full_report_{run_slug}.xlsx"

    print(
        f"Copy progress: preparing {total_tasks} tasks from {len(files)} source files"
    )
    print(f"Execution progress: 0/{total_tasks} pass=0 fail=0 ETA=? jobs={jobs}")

    tasks: list[tuple[_Case, Path | None]] = []
    for case in CASES:
        if case.uses_files:
            tasks.extend((case, source_file) for source_file in files)
        else:
            tasks.append((case, None))

    started = time.perf_counter()
    completed = 0
    pass_count = 0
    fail_count = 0
    failure_snapshots: list[dict[str, Any]] = []
    results: list[dict[str, Any]] = []

    with progress_path.open("w", encoding="utf-8", newline="\n") as progress_handle:
        _write_progress_event(
            progress_handle,
            {
                "event": "start",
                "timestamp_utc": datetime.now(tz=UTC).isoformat(),
                "run_root": str(run_root),
                "total_tasks": total_tasks,
                "jobs": jobs,
            },
        )

        with ThreadPoolExecutor(max_workers=jobs) as executor:
            future_map = {
                executor.submit(
                    _run_case_for_file,
                    repo_root=repo_root,
                    run_root=run_root,
                    oracle_exe=args.oracle_exe,
                    case=case,
                    source_file=source_file,
                ): (case, source_file)
                for case, source_file in tasks
            }

            for _future, (case, source_file) in future_map.items():
                _write_progress_event(
                    progress_handle,
                    {
                        "event": "scheduled",
                        "task_id": _file_task_id(case, source_file),
                        "case": case.name,
                        "file": source_file.name if source_file is not None else None,
                    },
                )

            for future in as_completed(future_map):
                result = future.result()
                results.append(result)
                completed += 1
                if result["case_pass"]:
                    pass_count += 1
                else:
                    fail_count += 1
                    failure_snapshots.append(
                        {
                            "task_id": result["task_id"],
                            "case": result["case"],
                            "file": result["file"],
                            "gate_checks": result["gate_checks"],
                        }
                    )
                _write_progress_event(
                    progress_handle,
                    {
                        "event": "completed",
                        "task_id": result["task_id"],
                        "case": result["case"],
                        "file": result["file"],
                        "case_pass": result["case_pass"],
                        "completed": completed,
                        "total": total_tasks,
                        "pass_count": pass_count,
                        "fail_count": fail_count,
                    },
                )
                print(
                    "Execution progress: "
                    f"{completed}/{total_tasks} pass={pass_count} fail={fail_count} "
                    f"ETA={_eta_text(completed=completed, total=total_tasks, started=started)}"
                )
                if fail_count > 0 and fail_count % 5 == 0:
                    snapshot = failure_snapshots[-3:]
                    print(
                        f"Failure snapshot ({fail_count} fails): {json.dumps(snapshot, sort_keys=True)}"
                    )
                    _write_progress_event(
                        progress_handle,
                        {
                            "event": "failure_snapshot",
                            "fail_count": fail_count,
                            "snapshot": snapshot,
                        },
                    )

        _write_progress_event(
            progress_handle,
            {
                "event": "finished",
                "timestamp_utc": datetime.now(tz=UTC).isoformat(),
                "completed": completed,
                "pass_count": pass_count,
                "fail_count": fail_count,
            },
        )

    results.sort(key=lambda item: (item["case"], item["file"] or ""))
    source_hash_after = {path.name: sha256_file(path) for path in files}
    source_non_mutated = source_hash_before == source_hash_after

    gating_results = [item for item in results if item["gating"]]
    gating_pass = sum(1 for item in gating_results if item["case_pass"])
    gating_fail = len(gating_results) - gating_pass

    per_case: dict[str, dict[str, Any]] = {}
    for item in results:
        row = per_case.setdefault(
            item["case"],
            {
                "case": item["case"],
                "runs": 0,
                "pass": 0,
                "fail": 0,
                "gating": item["gating"],
            },
        )
        row["runs"] += 1
        if item["case_pass"]:
            row["pass"] += 1
        else:
            row["fail"] += 1

    case_rows = [per_case[name] for name in sorted(per_case)]
    failure_rows = [item for item in results if not item["case_pass"]]

    summary = {
        "status": "PASS" if gating_fail == 0 and source_non_mutated else "FAIL",
        "required_gating_pass": gating_pass,
        "required_gating_fail": gating_fail,
        "source_non_mutated": source_non_mutated,
        "files": len(files),
        "total_tasks": total_tasks,
        "jobs_used": jobs,
        "run_root": str(run_root),
        "source_dir": str(args.source_dir),
        "oracle_exe": str(args.oracle_exe),
        "snapshot_dir": str(args.snapshot_dir),
    }

    results_payload = {
        "summary": summary,
        "results": results,
    }
    results_path.write_text(
        json.dumps(results_payload, indent=2, sort_keys=True), encoding="utf-8"
    )

    run_payload = {
        "summary": summary,
        "gates": {
            "required_cases_pass": gating_fail == 0,
            "source_non_mutated": source_non_mutated,
        },
        "artifacts": {
            "results_json": str(results_path),
            "progress_jsonl": str(progress_path),
            "report_xlsx": str(report_xlsx),
            "run_json": str(run_path),
        },
        "source_hash_before": source_hash_before,
        "source_hash_after": source_hash_after,
    }
    run_path.write_text(
        json.dumps(run_payload, indent=2, sort_keys=True), encoding="utf-8"
    )

    _write_xlsx(
        summary=summary,
        case_rows=case_rows,
        per_file_rows=results,
        failure_rows=failure_rows,
        path=report_xlsx,
    )

    print(f"Run artifact: {run_path}")
    print(f"Results artifact: {results_path}")
    print(f"Progress artifact: {progress_path}")
    print(f"Report artifact: {report_xlsx}")
    return 0 if summary["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
