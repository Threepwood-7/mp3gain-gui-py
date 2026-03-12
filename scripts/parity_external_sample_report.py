"""Run strict parity on a seeded external 30-file sample and export Excel report."""

from __future__ import annotations

import argparse
import json
import os
import random
import shutil
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from parity_common import (
    CODEX_89DB_DIR,
    default_parity_jobs,
    list_mp3_files,
    sha256_file,
)

if TYPE_CHECKING:
    from argparse import Namespace

    from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_EXTERNAL_ORIGINAL_DIR = Path("mp3-albums/original")
DEFAULT_EXTERNAL_REFERENCE_DIR = Path("mp3-albums/reference-normalized")
DEFAULT_REPORT_DIR = Path("c:/tmp/pycompa")
DEFAULT_SAMPLE_SIZE = 30
DEFAULT_SAMPLE_SEED = 42
DEFAULT_SPOT_CHECK_COUNT = 3
DEFAULT_JOBS = min(8, max(1, os.cpu_count() or 1))
LEGACY_PROFILES = ("skip_tags", "write_apev2", "write_id3")

PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")


def _parse_args() -> Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--original-dir",
        type=Path,
        default=DEFAULT_EXTERNAL_ORIGINAL_DIR,
        help="External originals directory.",
    )
    parser.add_argument(
        "--reference-dir",
        type=Path,
        default=DEFAULT_EXTERNAL_REFERENCE_DIR,
        help="External legacy 89 dB directory.",
    )
    parser.add_argument(
        "--workspace-root",
        type=Path,
        default=CODEX_89DB_DIR,
        help="Workspace root where temporary sampled copies are created.",
    )
    parser.add_argument(
        "--snapshot-dir",
        type=Path,
        default=DEFAULT_REPORT_DIR,
        help="Directory where timestamped JSON/XLSX artifacts are written.",
    )
    parser.add_argument(
        "--sample-size",
        type=int,
        default=DEFAULT_SAMPLE_SIZE,
        help="Number of files to sample from filename intersection.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=DEFAULT_SAMPLE_SEED,
        help="Random seed used for deterministic sample selection.",
    )
    parser.add_argument(
        "--jobs",
        type=int,
        default=max(DEFAULT_JOBS, default_parity_jobs()),
        help="Parallel worker count for copy and parity run.",
    )
    parser.add_argument(
        "--strict",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Use strict parity gates (default: true).",
    )
    parser.add_argument(
        "--force-rebuild",
        action="store_true",
        help="Pass --force-rebuild into parity build step.",
    )
    parser.add_argument(
        "--legacy-profile",
        choices=LEGACY_PROFILES,
        default="skip_tags",
        help="Legacy profile passed through to parity_compare/parity_build.",
    )
    return parser.parse_args()


def _status(value: bool) -> str:
    return "PASS" if value else "FAIL"


def _normalize_file_map(files: list[Path]) -> dict[str, Path]:
    return {path.name: path for path in files}


def _select_sample_names(
    common_names: list[str], *, sample_size: int, seed: int
) -> list[str]:
    selected = random.Random(seed).sample(common_names, sample_size)
    return sorted(selected, key=str.lower)


def _spot_hashes(
    file_map: dict[str, Path], sample_names: list[str], count: int
) -> dict[str, str]:
    names = sample_names[:count]
    return {name: sha256_file(file_map[name]) for name in names}


def _copy_one(src: Path, dst: Path) -> tuple[str, str]:
    shutil.copyfile(src, dst)
    return src.name, sha256_file(dst)


def _parallel_copy(tasks: list[tuple[Path, Path]], jobs: int) -> dict[str, str]:
    copied_hashes: dict[str, str] = {}
    if not tasks:
        return copied_hashes

    max_workers = min(max(1, jobs), len(tasks))
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        future_map = {
            executor.submit(_copy_one, src, dst): src.name for src, dst in tasks
        }
        for future in as_completed(future_map):
            name, copied_hash = future.result()
            copied_hashes[name] = copied_hash
    return copied_hashes


def _ensure_clean_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def _sheet_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _autosize(ws: Worksheet) -> None:
    for column in ws.columns:
        column_cells = list(column)
        if not column_cells:
            continue
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = min(80, max(12, max_len + 2))


def _apply_status_fills(ws: Worksheet) -> None:
    for row in ws.iter_rows(
        min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column
    ):
        for cell in row:
            if cell.value == "PASS":
                cell.fill = PASS_FILL
            elif cell.value == "FAIL":
                cell.fill = FAIL_FILL


def _render_mismatch(mismatch: object) -> str:
    if not isinstance(mismatch, dict):
        return ""
    reason = str(mismatch.get("reason", "")).strip()
    if reason:
        return reason
    if "offset" in mismatch:
        return f"offset={mismatch['offset']}"
    return json.dumps(mismatch, sort_keys=True)


def _flatten_tag_rows(tag_mismatches: object) -> list[tuple[str, str, str, str, str]]:
    rows: list[tuple[str, str, str, str, str]] = []
    if not isinstance(tag_mismatches, dict):
        return rows
    for filename in sorted(tag_mismatches):
        item = tag_mismatches.get(filename)
        if not isinstance(item, dict):
            continue
        missing_in_codex = item.get("missing_in_codex")
        if isinstance(missing_in_codex, list):
            for key in missing_in_codex:
                rows.append((str(filename), "missing_in_codex", str(key), "", ""))
        missing_in_reference = item.get("missing_in_reference")
        if isinstance(missing_in_reference, list):
            for key in missing_in_reference:
                rows.append((str(filename), "missing_in_reference", str(key), "", ""))
        value_diffs = item.get("value_diffs")
        if isinstance(value_diffs, dict):
            for key in sorted(value_diffs):
                diff = value_diffs.get(key)
                if not isinstance(diff, dict):
                    continue
                rows.append(
                    (
                        str(filename),
                        "value_diff",
                        str(key),
                        str(diff.get("reference", "")),
                        str(diff.get("codex", "")),
                    )
                )
    return rows


def _flatten_table_rows(
    table_mismatches: object,
) -> list[tuple[str, str, str, str, str]]:
    rows: list[tuple[str, str, str, str, str]] = []
    if not isinstance(table_mismatches, dict):
        return rows
    for row_key in sorted(table_mismatches):
        item = table_mismatches.get(row_key)
        if not isinstance(item, dict):
            continue
        missing_in_codex = item.get("missing_in_codex")
        if isinstance(missing_in_codex, dict):
            for field in sorted(missing_in_codex):
                rows.append(
                    (
                        str(row_key),
                        "missing_in_codex",
                        str(field),
                        str(missing_in_codex.get(field, "")),
                        "",
                    )
                )
        elif isinstance(missing_in_codex, list):
            for field in missing_in_codex:
                rows.append((str(row_key), "missing_in_codex", str(field), "", ""))

        missing_in_reference = item.get("missing_in_reference")
        if isinstance(missing_in_reference, dict):
            for field in sorted(missing_in_reference):
                rows.append(
                    (
                        str(row_key),
                        "missing_in_reference",
                        str(field),
                        "",
                        str(missing_in_reference.get(field, "")),
                    )
                )
        elif isinstance(missing_in_reference, list):
            for field in missing_in_reference:
                rows.append((str(row_key), "missing_in_reference", str(field), "", ""))

        value_diffs = item.get("value_diffs")
        if isinstance(value_diffs, dict):
            for field in sorted(value_diffs):
                diff = value_diffs.get(field)
                if not isinstance(diff, dict):
                    continue
                rows.append(
                    (
                        str(row_key),
                        "value_diff",
                        str(field),
                        str(diff.get("reference", "")),
                        str(diff.get("codex", "")),
                    )
                )
    return rows


def _write_excel_report(
    *,
    compare_report: dict[str, object],
    run_report: dict[str, object],
    output_path: Path,
) -> None:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    _sheet_header(summary_ws, ["Metric", "Value"])

    summary_rows = [
        ("created_at_utc", run_report.get("created_at_utc", "")),
        ("strict", _status(bool(run_report.get("strict", False)))),
        ("legacy_profile", run_report.get("legacy_profile", "")),
        ("compare_exit_code", run_report.get("compare_exit_code", "")),
        ("compare_status", str(compare_report.get("summary", {}).get("status", ""))),
        ("sample_size", run_report.get("sample_size", "")),
        ("seed", run_report.get("seed", "")),
        ("jobs_requested", run_report.get("jobs_requested", "")),
        ("jobs_copy_used", run_report.get("jobs_copy_used", "")),
        ("jobs_parity_used", compare_report.get("summary", {}).get("jobs_used", "")),
        ("source_original_dir", run_report.get("source_original_dir", "")),
        ("source_reference_dir", run_report.get("source_reference_dir", "")),
        ("workspace_root", run_report.get("workspace_root", "")),
        ("workspace_original_dir", run_report.get("workspace_original_dir", "")),
        ("workspace_reference_dir", run_report.get("workspace_reference_dir", "")),
        ("workspace_codex_dir", run_report.get("workspace_codex_dir", "")),
        ("workspace_snapshot_dir", run_report.get("workspace_snapshot_dir", "")),
        ("selected_common_files", len(run_report.get("selected_files", []))),
        ("reference_extra_files", len(run_report.get("extra_in_reference", []))),
        ("original_missing_files", len(run_report.get("missing_in_reference", []))),
        (
            "source_files_unchanged",
            _status(bool(run_report.get("spot_check_unchanged", False))),
        ),
    ]
    for key, value in summary_rows:
        summary_ws.append([key, value])

    statuses_obj = compare_report.get("statuses")
    if isinstance(statuses_obj, dict):
        for status_key in sorted(statuses_obj):
            summary_ws.append([f"status.{status_key}", statuses_obj.get(status_key)])

    _apply_status_fills(summary_ws)
    _autosize(summary_ws)

    per_file_ws = wb.create_sheet("PerFile")
    per_file_headers = [
        "file",
        "size_match",
        "hash_match",
        "global_gain_ok",
        "container_ok",
        "id3v1_match",
        "reference_size",
        "codex_size",
        "first_global_gain_mismatch",
        "first_byte_mismatch",
        "byte_mismatch_bucket",
    ]
    _sheet_header(per_file_ws, per_file_headers)

    files_obj = compare_report.get("files")
    if isinstance(files_obj, dict):
        for filename in sorted(files_obj):
            item = files_obj.get(filename)
            if not isinstance(item, dict):
                continue
            byte_mismatch = item.get("first_byte_mismatch")
            bucket = ""
            if isinstance(byte_mismatch, dict):
                bucket = str(byte_mismatch.get("semantic_bucket", ""))
            per_file_ws.append(
                [
                    filename,
                    _status(bool(item.get("size_match", False))),
                    _status(bool(item.get("hash_match", False))),
                    _status(
                        bool(
                            item.get(
                                "global_gain_ok",
                                "first_global_gain_mismatch" not in item,
                            )
                        )
                    ),
                    _status(
                        bool(item.get("container_ok", item.get("id3v1_match", False)))
                    ),
                    _status(bool(item.get("id3v1_match", False))),
                    item.get("reference_size", ""),
                    item.get("codex_size", ""),
                    _render_mismatch(item.get("first_global_gain_mismatch")),
                    _render_mismatch(byte_mismatch),
                    bucket,
                ]
            )

    _apply_status_fills(per_file_ws)
    _autosize(per_file_ws)

    tag_ws = wb.create_sheet("TagDiff")
    _sheet_header(tag_ws, ["file", "diff_type", "key", "reference", "codex"])
    for row in _flatten_tag_rows(compare_report.get("tag_mismatches")):
        tag_ws.append(list(row))
    _autosize(tag_ws)

    table_ws = wb.create_sheet("TableDiff")
    _sheet_header(table_ws, ["row_key", "diff_type", "field", "reference", "codex"])
    for row in _flatten_table_rows(compare_report.get("table_mismatches")):
        table_ws.append(list(row))
    _autosize(table_ws)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> int:
    args = _parse_args()

    sample_size = max(1, int(args.sample_size))
    seed = int(args.seed)
    jobs_requested = max(1, int(args.jobs))
    strict = bool(args.strict)
    legacy_profile = str(args.legacy_profile)

    source_original_dir = args.original_dir
    source_reference_dir = args.reference_dir
    workspace_root = args.workspace_root
    report_dir = args.snapshot_dir

    sample_workspace = workspace_root / f"external_sample_n{sample_size}_seed{seed}"
    workspace_original_dir = sample_workspace / "original"
    workspace_reference_dir = sample_workspace / "reference_89db"
    workspace_codex_dir = sample_workspace / "codex_89db"
    workspace_snapshot_dir = sample_workspace / "snapshots"

    repo_root = Path(__file__).resolve().parents[1]

    source_original_map = _normalize_file_map(list_mp3_files(source_original_dir))
    source_reference_map = _normalize_file_map(list_mp3_files(source_reference_dir))
    source_original_names = set(source_original_map)
    source_reference_names = set(source_reference_map)

    common_names = sorted(source_original_names & source_reference_names, key=str.lower)
    missing_in_reference = sorted(
        source_original_names - source_reference_names, key=str.lower
    )
    extra_in_reference = sorted(
        source_reference_names - source_original_names, key=str.lower
    )

    if len(common_names) < sample_size:
        raise RuntimeError(
            f"Not enough common files for sample_size={sample_size}: "
            f"common={len(common_names)}"
        )

    selected_names = _select_sample_names(
        common_names, sample_size=sample_size, seed=seed
    )

    pre_spot_original = _spot_hashes(
        source_original_map, selected_names, DEFAULT_SPOT_CHECK_COUNT
    )
    pre_spot_reference = _spot_hashes(
        source_reference_map, selected_names, DEFAULT_SPOT_CHECK_COUNT
    )

    _ensure_clean_dir(sample_workspace)
    workspace_original_dir.mkdir(parents=True, exist_ok=True)
    workspace_reference_dir.mkdir(parents=True, exist_ok=True)
    workspace_codex_dir.mkdir(parents=True, exist_ok=True)
    workspace_snapshot_dir.mkdir(parents=True, exist_ok=True)

    copy_tasks: list[tuple[Path, Path]] = []
    for name in selected_names:
        copy_tasks.append((source_original_map[name], workspace_original_dir / name))
        copy_tasks.append((source_reference_map[name], workspace_reference_dir / name))

    jobs_copy_used = min(max(1, jobs_requested), len(copy_tasks))
    copied_hashes = _parallel_copy(copy_tasks, jobs_copy_used)

    compare_command = [
        sys.executable,
        "scripts/parity_compare.py",
        "--mode",
        "full",
        "--jobs",
        str(jobs_requested),
        "--original-dir",
        str(workspace_original_dir),
        "--reference-dir",
        str(workspace_reference_dir),
        "--codex-dir",
        str(workspace_codex_dir),
        "--snapshot-dir",
        str(workspace_snapshot_dir),
        "--legacy-profile",
        legacy_profile,
    ]
    if strict:
        compare_command.append("--strict")
    if args.force_rebuild:
        compare_command.append("--force-rebuild")

    completed = subprocess.run(
        compare_command,
        cwd=repo_root,
        check=False,
        text=True,
        capture_output=True,
    )

    compare_report_path = workspace_snapshot_dir / "parity_compare_report.json"
    if not compare_report_path.exists():
        raise RuntimeError(
            "Parity compare report was not created. "
            f"exit_code={completed.returncode}, stderr={completed.stderr}"
        )
    compare_report = json.loads(compare_report_path.read_text(encoding="utf-8"))
    if not isinstance(compare_report, dict):
        raise RuntimeError("parity_compare_report.json is not an object.")

    post_spot_original = _spot_hashes(
        source_original_map, selected_names, DEFAULT_SPOT_CHECK_COUNT
    )
    post_spot_reference = _spot_hashes(
        source_reference_map, selected_names, DEFAULT_SPOT_CHECK_COUNT
    )
    spot_check_unchanged = (
        pre_spot_original == post_spot_original
        and pre_spot_reference == post_spot_reference
    )

    created_at = datetime.now(UTC)
    created_at_utc = created_at.isoformat()
    timestamp = created_at.strftime("%Y%m%d_%H%M%S")

    report_dir.mkdir(parents=True, exist_ok=True)
    external_compare_path = report_dir / f"parity_external_compare_{timestamp}.json"
    external_run_path = report_dir / f"parity_external_run_{timestamp}.json"
    external_excel_path = report_dir / f"parity_external_report_{timestamp}.xlsx"

    external_compare_path.write_text(
        json.dumps(compare_report, indent=2, sort_keys=True),
        encoding="utf-8",
    )

    run_report: dict[str, object] = {
        "created_at_utc": created_at_utc,
        "sample_size": sample_size,
        "seed": seed,
        "strict": strict,
        "jobs_requested": jobs_requested,
        "jobs_copy_used": jobs_copy_used,
        "compare_exit_code": completed.returncode,
        "legacy_profile": legacy_profile,
        "compare_command": compare_command,
        "compare_stdout": completed.stdout,
        "compare_stderr": completed.stderr,
        "source_original_dir": str(source_original_dir),
        "source_reference_dir": str(source_reference_dir),
        "workspace_root": str(workspace_root),
        "workspace_original_dir": str(workspace_original_dir),
        "workspace_reference_dir": str(workspace_reference_dir),
        "workspace_codex_dir": str(workspace_codex_dir),
        "workspace_snapshot_dir": str(workspace_snapshot_dir),
        "selected_files": selected_names,
        "common_count": len(common_names),
        "missing_in_reference": missing_in_reference,
        "extra_in_reference": extra_in_reference,
        "source_spot_check_pre_original_sha256": pre_spot_original,
        "source_spot_check_post_original_sha256": post_spot_original,
        "source_spot_check_pre_reference_sha256": pre_spot_reference,
        "source_spot_check_post_reference_sha256": post_spot_reference,
        "spot_check_unchanged": spot_check_unchanged,
        "copied_sha256": copied_hashes,
        "artifacts": {
            "workspace_compare_report": str(compare_report_path),
            "external_compare_json": str(external_compare_path),
            "external_run_json": str(external_run_path),
            "external_excel_report": str(external_excel_path),
        },
    }

    external_run_path.write_text(
        json.dumps(run_report, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    _write_excel_report(
        compare_report=compare_report,
        run_report=run_report,
        output_path=external_excel_path,
    )

    print(f"Sample workspace prepared: {sample_workspace}")
    print(
        f"Common files={len(common_names)} selected={len(selected_names)} seed={seed}"
    )
    print(f"Extra in reference (ignored for sampling): {len(extra_in_reference)}")
    print(f"Original/source spot-check unchanged: {_status(spot_check_unchanged)}")
    print(f"Run report: {external_run_path}")
    print(f"Compare report: {external_compare_path}")
    print(f"Excel report: {external_excel_path}")

    return completed.returncode


if __name__ == "__main__":
    raise SystemExit(main())
